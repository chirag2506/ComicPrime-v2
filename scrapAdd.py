from bs4 import BeautifulSoup as soup
from urllib.request import urlopen as uReq
from datetime import date
from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.shared import Inches
import calendar

DOCUMENT_NAME = 'ComicsNew.docx'
# DOCUMENT_NAME = 'Comics to read (1939-Feb 2021).docx'
# DOCUMENT_NAME = 'Comics  (1939-Feb 2021).docx'
URL = 'https://marvel.fandom.com/wiki/Western_Gunfighters_Vol_2'
URL = 'https://marvel.fandom.com/wiki/Marvel_Spotlight_Vol_1'

document = Document(DOCUMENT_NAME)

paras = []
for paragraph in document.paragraphs:
        paras.append(paragraph.text)

my_url = URL
uClient = uReq(my_url)
page_html = uClient.read()
uClient.close()
page_soup = soup(page_html, "html.parser")

containers = page_soup.findAll("div", {"class": "wikia-gallery-item"})
comics_in_this_volume = len(containers)
print("COMICS IN THIS VOLUME: " + str(comics_in_this_volume))

for comic in range(0, len(containers)):
    container = containers[comic]
    comic_name = container.div.next_sibling.div.span.a.text
    release_and_cover = container.div.next_sibling.div.span.next_sibling.text
    cover_date = release_and_cover[release_and_cover.find("Cover date: ")+12 : ]
    index_of_cover = paras.index(cover_date)
    #changing number of comics
    number_in_month = document.paragraphs[index_of_cover + 1 + comic]
    number_in_month_in_new_text = int(number_in_month.text[38:]) + 1
    number_in_month.text = "Number of comics published this month: "
    number_in_month.add_run(str(number_in_month_in_new_text)).bold=True
    #adding the comic
    document.paragraphs[index_of_cover+2+comic].insert_paragraph_before(
        comic_name.strip(), style='List Bullet'
    )

    document.save(DOCUMENT_NAME)

document.save(DOCUMENT_NAME)